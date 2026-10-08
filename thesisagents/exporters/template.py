"""The PPTX template contract: what a custom template must offer and may override.

``ExportOptions.pptx_template`` used to mean "open this file and hope": the
exporter took ``slide_layouts[6]`` whatever it was, so a template with five
layouts crashed with ``list index out of range`` and one whose seventh layout
had placeholders shipped "Click to add title" boxes on every slide. This
module replaces that with a documented contract and a check that runs before
anything is rendered.

The contract (version 1)
------------------------
**Slide size.** The exporter lays out for 16:9 (13.333 x 7.5 in). A template
of another size is resized to 16:9 and reported (``slide_size = "normalise"``,
the default) or refused (``slide_size = "reject"``).

**Layouts by role.** Every slide the exporter makes has one of six roles
(:data:`ROLES`). A role's layout is, in this order: the layout the config
names for it, a layout named after the role, the ``content`` layout, and for
``content`` itself the template's blank layout. Artwork on the layout and its
master (background, logo, decorative shapes) shows through. Placeholders the
layout would put on the slide are removed, because the exporter draws its own
named text boxes at positions the overflow check knows.

**Title placeholder (opt-in).** For the roles in
:data:`TITLE_PLACEHOLDER_ROLES`, the config can ask for the slide title to be
written into the layout's title placeholder, so it sits where the template
puts titles. The layout must then have a title placeholder that ends above
the content area, and the exporter does not draw its header band there.

**Style overrides.** A config (TOML or JSON) may set the font families, the
four palette colours and whether the header band and the cover panel are
drawn. Font sizes and margins are fixed: slide geometry and the overflow
check are calibrated for the built-in values.

Example config (``thesis.toml``)::

    slide_size = "reject"

    [layouts]
    cover = "Title Slide"
    content = "Title Only"

    [placeholders]
    title = ["content", "table", "references"]

    [fonts]
    latin = "Source Sans 3"
    east_asian = "Noto Sans TC"

    [colors]
    primary = "#0B3D2E"
    highlight = "#0E7C61"

    [chrome]
    header_band = false

Example use::

    report = validate_template("thesis.pptx", "thesis.toml")
    if not report.ok:
        print(report.message())
"""

from __future__ import annotations

import json
import tomllib
from collections.abc import Mapping
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.util import Inches

from thesisagents.core.constants import EXPORT_PPTX
from thesisagents.core.exceptions import ExportError

#: Every slide the exporter makes has one of these roles.
ROLES: tuple[str, ...] = ("cover", "section", "content", "table", "references", "qa")
ROLE_COVER, ROLE_SECTION, ROLE_CONTENT, ROLE_TABLE, ROLE_REFERENCES, ROLE_QA = ROLES
#: Roles whose title can go into the layout's title placeholder: the ones
#: built as "title on top, content below". The cover, section and Q&A slides
#: place several centred text boxes and keep drawing their own.
TITLE_PLACEHOLDER_ROLES: tuple[str, ...] = (ROLE_CONTENT, ROLE_TABLE, ROLE_REFERENCES)
#: Palette colours a config may override, by what they are used for.
COLOR_KEYS: tuple[str, ...] = ("primary", "highlight", "muted", "subtle")

SIZE_NORMALISE = "normalise"
SIZE_REJECT = "reject"

#: The canvas every position in ``pptx.py`` is computed for.
SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)
#: Where the title area of a slide ends: the line under the title (the paper
#: subtitle, placed by ``pptx._add_paper_subtitle`` at this very constant)
#: starts here. A title placeholder must end above it or that line is drawn
#: over the title. The built-in header band ends at 1.24 in.
TITLE_AREA_BOTTOM = Inches(1.4)
_SIZE_TOLERANCE = Inches(0.01)
_EMU_PER_INCH = 914400

#: Layout names that mean a role when the config does not name one.
_ROLE_ALIASES: Mapping[str, tuple[str, ...]] = {
    ROLE_QA: ("qa", "q&a"),
    ROLE_REFERENCES: ("references", "reference"),
}
_BLANK_LAYOUT_NAME = "blank"
_TITLE_PLACEHOLDER_TYPES = (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)

#: A palette colour is text on a white slide, and ``primary`` is also the
#: fill behind white titles and table headers. Above this luminance (the
#: project's light-on-light threshold) either use becomes unreadable.
_MAX_LUMINANCE = 0.7 * 255
#: The warm red the project bans as a text colour (deck-design "No red text").
_BANNED_TEXT_COLOUR = (0xC0, 0x39, 0x2B)

_SECTIONS = ("slide_size", "layouts", "placeholders", "fonts", "colors", "chrome")
#: Sections people reasonably try, with why they are not offered.
_UNSUPPORTED_SECTIONS: Mapping[str, str] = {
    "margins": "margins are fixed",
    "sizes": "font sizes are fixed",
    "styles": "font sizes are fixed (families and colours go in [fonts] and [colors])",
}
_FIXED_GEOMETRY_REASON = (
    "slide geometry and the overflow check are calibrated for the built-in values"
)


class TemplateError(ExportError):
    """A template or its config does not meet the contract.

    Raised before any slide is rendered, with every problem found, so a long
    export does not fail at the end over a layout name.
    """

    def __init__(self, message: str) -> None:
        super().__init__(EXPORT_PPTX, message)


@dataclass(frozen=True, slots=True)
class TemplateConfig:
    """What a config file asks for. The default is "change nothing".

    ``layouts`` maps a role to a layout name. ``title_placeholder_roles`` are
    the roles whose title goes into the layout's title placeholder.
    ``colors`` maps a :data:`COLOR_KEYS` name to an ``(r, g, b)`` tuple.
    """

    layouts: Mapping[str, str] = field(default_factory=dict)
    title_placeholder_roles: frozenset[str] = frozenset()
    font_latin: str | None = None
    font_east_asian: str | None = None
    colors: Mapping[str, tuple[int, int, int]] = field(default_factory=dict)
    header_band: bool = True
    cover_panel: bool = True
    slide_size: str = SIZE_NORMALISE
    source: str | None = None


@dataclass(frozen=True, slots=True)
class TemplateIssue:
    """One finding. ``code`` is stable for scripts, ``message`` says what to do."""

    code: str
    message: str

    def to_dict(self) -> dict[str, str]:
        return {"code": self.code, "message": self.message}


@dataclass(frozen=True, slots=True)
class TemplateValidationReport:
    """What :func:`validate_template` found.

    ``layouts`` is the layout each role resolved to, so a user can see what an
    export would use before running one. ``errors`` stop an export,
    ``warnings`` do not.
    """

    template: str
    config: str | None = None
    slide_width_in: float | None = None
    slide_height_in: float | None = None
    available_layouts: tuple[str, ...] = ()
    layouts: Mapping[str, str] = field(default_factory=dict)
    errors: tuple[TemplateIssue, ...] = ()
    warnings: tuple[TemplateIssue, ...] = ()

    @property
    def ok(self) -> bool:
        return not self.errors

    def to_dict(self) -> dict[str, Any]:
        return {
            "template": self.template,
            "config": self.config,
            "ok": self.ok,
            "slide_width_in": self.slide_width_in,
            "slide_height_in": self.slide_height_in,
            "available_layouts": list(self.available_layouts),
            "layouts": dict(self.layouts),
            "errors": [issue.to_dict() for issue in self.errors],
            "warnings": [issue.to_dict() for issue in self.warnings],
        }

    def message(self) -> str:
        """Every error, one per line, for an exception or a terminal."""
        lines = [f"template {self.template} does not meet the template contract:"]
        lines.extend(f"  - {issue.message}" for issue in self.errors)
        return "\n".join(lines)


# ---------------------------------------------------------------------------
# Config
# ---------------------------------------------------------------------------


def load_template_config(path: str | Path) -> TemplateConfig:
    """Read a template config from a ``.toml`` or ``.json`` file.

    The boundary this guards: a file a person wrote by hand. Every problem in
    it is collected and raised together as one :class:`TemplateError`, each
    naming the key, so a config is fixed in one pass and not one error per
    run.

    Example: ``load_template_config("thesis.toml").colors["primary"]`` is
    ``(11, 61, 46)`` for ``primary = "#0B3D2E"``.
    """
    file = Path(path).expanduser()
    suffix = file.suffix.lower()
    if suffix not in {".toml", ".json"}:
        raise TemplateError(
            f"template config {file} must be a .toml or .json file"
        )
    try:
        raw = file.read_bytes()
    except OSError as err:
        raise TemplateError(f"cannot read template config {file}: {err}") from err
    try:
        data = tomllib.loads(raw.decode("utf-8")) if suffix == ".toml" else json.loads(raw)
    except (tomllib.TOMLDecodeError, json.JSONDecodeError, UnicodeDecodeError) as err:
        kind = suffix[1:].upper()
        raise TemplateError(f"template config {file} is not valid {kind}: {err}") from err
    if not isinstance(data, dict):
        raise TemplateError(f"template config {file} must hold a table of settings")
    return parse_template_config(data, source=str(file))


def parse_template_config(data: Mapping[str, Any], *, source: str | None = None) -> TemplateConfig:
    """Turn the decoded config into a :class:`TemplateConfig`, or say what is wrong."""
    problems: list[str] = []
    for key in data:
        if key in _UNSUPPORTED_SECTIONS:
            problems.append(
                f"[{key}] is not supported: {_UNSUPPORTED_SECTIONS[key]}, because "
                f"{_FIXED_GEOMETRY_REASON}"
            )
        elif key not in _SECTIONS:
            problems.append(f"unknown setting '{key}' (known: {', '.join(_SECTIONS)})")
    slide_size = data.get("slide_size", SIZE_NORMALISE)
    if slide_size not in (SIZE_NORMALISE, SIZE_REJECT):
        problems.append(
            f"slide_size must be \"{SIZE_NORMALISE}\" or \"{SIZE_REJECT}\", not {slide_size!r}"
        )
    config = TemplateConfig(
        layouts=_parse_layouts(_section(data, "layouts", problems), problems),
        title_placeholder_roles=_parse_title_roles(
            _section(data, "placeholders", problems), problems
        ),
        font_latin=_parse_font(_section(data, "fonts", problems), "latin", problems),
        font_east_asian=_parse_font(_section(data, "fonts", problems), "east_asian", problems),
        colors=_parse_colors(_section(data, "colors", problems), problems),
        header_band=_parse_flag(_section(data, "chrome", problems), "header_band", problems),
        cover_panel=_parse_flag(_section(data, "chrome", problems), "cover_panel", problems),
        slide_size=slide_size if slide_size in (SIZE_NORMALISE, SIZE_REJECT) else SIZE_NORMALISE,
        source=source,
    )
    # A section that is read for two settings reports its problems twice.
    unique = list(dict.fromkeys(problems))
    if unique:
        where = f"template config {source}" if source else "template config"
        raise TemplateError(
            f"{where} has {len(unique)} problem(s):\n"
            + "\n".join(f"  - {problem}" for problem in unique)
        )
    return config


def _section(data: Mapping[str, Any], name: str, problems: list[str]) -> Mapping[str, Any]:
    value = data.get(name, {})
    if isinstance(value, Mapping):
        return value
    problems.append(f"[{name}] must be a table of settings")
    return {}


def _unknown_keys(
    section: Mapping[str, Any], name: str, known: tuple[str, ...], problems: list[str]
) -> None:
    for key in section:
        if key not in known:
            problems.append(f"[{name}] has no setting '{key}' (known: {', '.join(known)})")


def _parse_layouts(section: Mapping[str, Any], problems: list[str]) -> dict[str, str]:
    _unknown_keys(section, "layouts", ROLES, problems)
    layouts: dict[str, str] = {}
    for role in ROLES:
        name = section.get(role)
        if name is None:
            continue
        if not isinstance(name, str) or not name.strip():
            problems.append(f"[layouts] {role} must be a layout name")
            continue
        layouts[role] = name.strip()
    return layouts


def _parse_title_roles(section: Mapping[str, Any], problems: list[str]) -> frozenset[str]:
    _unknown_keys(section, "placeholders", ("title",), problems)
    roles = section.get("title", [])
    if not isinstance(roles, list) or not all(isinstance(role, str) for role in roles):
        problems.append("[placeholders] title must be a list of roles")
        return frozenset()
    for role in roles:
        if role not in TITLE_PLACEHOLDER_ROLES:
            problems.append(
                f"[placeholders] title cannot include '{role}': the title placeholder "
                f"is offered for {', '.join(TITLE_PLACEHOLDER_ROLES)}"
            )
    return frozenset(role for role in roles if role in TITLE_PLACEHOLDER_ROLES)


def _parse_font(section: Mapping[str, Any], key: str, problems: list[str]) -> str | None:
    _unknown_keys(section, "fonts", ("latin", "east_asian"), problems)
    value = section.get(key)
    if value is None:
        return None
    if not isinstance(value, str) or not value.strip():
        problems.append(f"[fonts] {key} must be a font family name")
        return None
    return value.strip()


def _parse_flag(section: Mapping[str, Any], key: str, problems: list[str]) -> bool:
    _unknown_keys(section, "chrome", ("header_band", "cover_panel"), problems)
    value = section.get(key, True)
    if not isinstance(value, bool):
        problems.append(f"[chrome] {key} must be true or false")
        return True
    return value


def _parse_colors(
    section: Mapping[str, Any], problems: list[str]
) -> dict[str, tuple[int, int, int]]:
    _unknown_keys(section, "colors", COLOR_KEYS, problems)
    colors: dict[str, tuple[int, int, int]] = {}
    for key in COLOR_KEYS:
        value = section.get(key)
        if value is None:
            continue
        rgb = _parse_hex(value)
        if rgb is None:
            problems.append(f"[colors] {key} must be a colour like \"#1F3A66\", not {value!r}")
            continue
        if rgb == _BANNED_TEXT_COLOUR:
            problems.append(
                f"[colors] {key} is #C0392B, the red this project bans as a text "
                "colour (it reads as an error or a warning on a slide). Pick another."
            )
            continue
        if _luminance(rgb) > _MAX_LUMINANCE:
            problems.append(
                f"[colors] {key} = {value} is too light: it is used as text on a "
                "white slide"
                + (", and as the fill behind white titles" if key == "primary" else "")
                + ". Pick a darker colour."
            )
            continue
        colors[key] = rgb
    return colors


def _parse_hex(value: object) -> tuple[int, int, int] | None:
    if not isinstance(value, str):
        return None
    text = value.strip().removeprefix("#")
    if len(text) != 6:
        return None
    try:
        return int(text[0:2], 16), int(text[2:4], 16), int(text[4:6], 16)
    except ValueError:
        return None


def _luminance(rgb: tuple[int, int, int]) -> float:
    red, green, blue = rgb
    return 0.299 * red + 0.587 * green + 0.114 * blue


# ---------------------------------------------------------------------------
# Layouts
# ---------------------------------------------------------------------------


class LayoutSet:
    """The layout for each role of one presentation, and how slides are added.

    The exporter's builders receive this where they used to receive the one
    blank layout, and say which role the slide they are adding has.

    Example::

        slide = layouts.add_slide(prs, ROLE_TABLE)
        placeholder = layouts.title_placeholder(slide, ROLE_TABLE)   # or None
    """

    def __init__(
        self, layouts: Mapping[str, Any], title_roles: frozenset[str] = frozenset()
    ) -> None:
        self._layouts = dict(layouts)
        self._title_roles = title_roles

    def layout(self, role: str) -> Any:
        return self._layouts[role]

    def add_slide(self, prs: Presentation, role: str) -> Any:
        """Add a slide for ``role`` and clear the placeholders it inherits.

        A layout puts a copy of each of its placeholders on a new slide. Left
        empty they show "Click to add title" in PowerPoint, so all are
        removed, except the title placeholder of a role that writes into it.
        """
        slide = prs.slides.add_slide(self._layouts[role])
        keep = _title_placeholder(slide) if role in self._title_roles else None
        for placeholder in list(slide.placeholders):
            if keep is not None and placeholder._element is keep._element:
                continue
            placeholder._element.getparent().remove(placeholder._element)
        return slide

    def title_placeholder(self, slide: Any, role: str) -> Any | None:
        """The slide's title placeholder when ``role`` writes its title there."""
        if role not in self._title_roles:
            return None
        return _title_placeholder(slide)


def _title_placeholder(holder: Any) -> Any | None:
    """The title placeholder of a slide or layout, or ``None``."""
    for placeholder in holder.placeholders:
        if placeholder.placeholder_format.type in _TITLE_PLACEHOLDER_TYPES:
            return placeholder
    return None


def builtin_layouts(prs: Presentation) -> LayoutSet:
    """Every role on the blank layout of python-pptx's own template.

    This is the built-in navy-band deck: no template artwork, everything
    drawn by the exporter.
    """
    blank = _blank_layout(prs)
    return LayoutSet(dict.fromkeys(ROLES, blank))


def content_layout(prs: Presentation) -> Any:
    """A layout to put a text slide on, for a deck that is already built.

    ``pptx_edit.add_slide`` used ``slide_layouts[6]``, which is the blank
    layout only in python-pptx's own template. On a deck made from a custom
    template that index can be anything, or missing.
    """
    layouts, errors, _ = _resolve_layouts(prs, TemplateConfig())
    if errors:
        return prs.slide_layouts[len(prs.slide_layouts) - 1]
    return layouts[ROLE_CONTENT]


def _blank_layout(prs: Presentation) -> Any | None:
    for layout in prs.slide_layouts:
        if layout.name.casefold() == _BLANK_LAYOUT_NAME:
            return layout
    for layout in prs.slide_layouts:
        if not list(layout.placeholders):
            return layout
    return None


def _find_layout(prs: Presentation, name: str) -> Any | None:
    wanted = name.casefold()
    for layout in prs.slide_layouts:
        if layout.name == name:
            return layout
    for layout in prs.slide_layouts:
        if layout.name.casefold() == wanted:
            return layout
    return None


def _resolve_layouts(
    prs: Presentation, config: TemplateConfig
) -> tuple[dict[str, Any], list[TemplateIssue], list[TemplateIssue]]:
    """Find the layout of every role. Returns ``(layouts, errors, warnings)``."""
    errors: list[TemplateIssue] = []
    warnings: list[TemplateIssue] = []
    available = ", ".join(f"'{layout.name}'" for layout in prs.slide_layouts) or "none"
    found = _layouts_by_name(prs, config, errors, available)
    if ROLE_CONTENT not in found and ROLE_CONTENT not in config.layouts:
        blank = _blank_layout(prs)
        if blank is None:
            errors.append(TemplateIssue(
                "no-content-layout",
                "the template has no layout for slide content. Add a layout named "
                "'content' or 'Blank' (PowerPoint: View > Slide Master), or name one "
                f"in the config with [layouts] content = \"...\". Its layouts are: {available}.",
            ))
        else:
            found[ROLE_CONTENT] = blank
    if ROLE_CONTENT in found:
        for role in ROLES:
            if role not in found and role not in config.layouts:
                found[role] = found[ROLE_CONTENT]
    for role in sorted(config.title_placeholder_roles):
        layout = found.get(role)
        if layout is not None:
            _check_title_placeholder(role, layout, errors)
    return found, errors, warnings


def _layouts_by_name(
    prs: Presentation, config: TemplateConfig, errors: list[TemplateIssue], available: str
) -> dict[str, Any]:
    """The roles whose layout is found by name: the config's, or the role's own."""
    found: dict[str, Any] = {}
    for role in ROLES:
        named = config.layouts.get(role)
        candidates = _ROLE_ALIASES.get(role, (role,)) if named is None else (named,)
        layout = next(
            (hit for hit in (_find_layout(prs, name) for name in candidates) if hit is not None),
            None,
        )
        if layout is not None:
            found[role] = layout
        elif named is not None:
            errors.append(TemplateIssue(
                "layout-not-found",
                f"[layouts] {role} = \"{named}\", and the template has no layout "
                f"with that name. Its layouts are: {available}.",
            ))
    return found


def _check_title_placeholder(role: str, layout: Any, errors: list[TemplateIssue]) -> None:
    placeholder = _title_placeholder(layout)
    if placeholder is None:
        errors.append(TemplateIssue(
            "missing-title-placeholder",
            f"[placeholders] title includes '{role}', and its layout "
            f"'{layout.name}' has no title placeholder. Add one to that layout "
            f"(PowerPoint: View > Slide Master), or remove '{role}' from "
            "[placeholders] title.",
        ))
        return
    top, height = placeholder.top, placeholder.height
    if top is None or height is None:
        return
    bottom = top + height
    limit = f"{TITLE_AREA_BOTTOM / _EMU_PER_INCH:.2f} in"
    if bottom > TITLE_AREA_BOTTOM:
        errors.append(TemplateIssue(
            "title-placeholder-overlaps-content",
            f"the title placeholder of layout '{layout.name}' (role '{role}') ends "
            f"{bottom / _EMU_PER_INCH:.2f} in from the top, and slide content starts at "
            f"{limit}, so the first line would be drawn over the title. Make the "
            f"placeholder end at or above {limit} (PowerPoint: View > Slide Master).",
        ))


# ---------------------------------------------------------------------------
# Validation and opening
# ---------------------------------------------------------------------------


@dataclass(frozen=True, slots=True)
class OpenedTemplate:
    """A presentation ready to be filled: the file, its layouts, its config."""

    presentation: Any
    layouts: LayoutSet
    config: TemplateConfig
    report: TemplateValidationReport


def validate_template(
    path: str | Path,
    config: TemplateConfig | str | Path | None = None,
    *,
    dark_mode: bool = False,
) -> TemplateValidationReport:
    """Check a template against the contract without rendering anything.

    The boundary this guards: the start of an export. A template problem
    found here costs a second. Found by the exporter after a search, PDF
    downloads and summarisation, it costs the whole run.

    ``config`` is a :class:`TemplateConfig` or the path of a config file. A
    config that cannot be read or parsed is reported as an error, like any
    other finding. ``dark_mode`` adds the warnings that apply to a dark export.

    Example::

        report = validate_template("thesis.pptx", "thesis.toml")
        report.ok            # False
        report.errors[0].code   # "layout-not-found"
    """
    return _inspect(path, config, dark_mode=dark_mode)[0]


def open_template(
    path: str | Path,
    config: TemplateConfig | str | Path | None = None,
    *,
    dark_mode: bool = False,
) -> OpenedTemplate:
    """Open a template for an export, or raise :class:`TemplateError`.

    Runs the same checks as :func:`validate_template` on the same
    presentation object the exporter then fills, so what was validated is
    what is rendered. The returned presentation is 16:9.
    """
    report, prs, layouts, parsed = _inspect(path, config, dark_mode=dark_mode)
    if not report.ok or prs is None or parsed is None:
        raise TemplateError(report.message())
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT
    return OpenedTemplate(
        presentation=prs,
        layouts=LayoutSet(layouts, parsed.title_placeholder_roles),
        config=parsed,
        report=report,
    )


def _inspect(
    path: str | Path,
    config: TemplateConfig | str | Path | None,
    *,
    dark_mode: bool,
) -> tuple[TemplateValidationReport, Any | None, dict[str, Any], TemplateConfig | None]:
    template = str(path)
    config_source = config.source if isinstance(config, TemplateConfig) else (
        str(config) if config is not None else None
    )
    errors: list[TemplateIssue] = []
    warnings: list[TemplateIssue] = []
    parsed: TemplateConfig | None
    if isinstance(config, TemplateConfig):
        parsed = config
    elif config is None:
        parsed = TemplateConfig()
    else:
        try:
            parsed = load_template_config(config)
        except TemplateError as err:
            parsed = None
            errors.append(TemplateIssue("invalid-config", str(err).removeprefix("[pptx] ")))
    try:
        prs = Presentation(template)
    except Exception as err:  # noqa: BLE001 - python-pptx raises several unrelated types
        errors.append(TemplateIssue(
            "unreadable-template",
            f"cannot open {template} as a PowerPoint file ({type(err).__name__}: {err}). "
            "A template is a .pptx or .potx saved by PowerPoint.",
        ))
        return (
            TemplateValidationReport(
                template=template, config=config_source,
                errors=tuple(errors), warnings=tuple(warnings),
            ),
            None, {}, parsed,
        )
    width, height = prs.slide_width, prs.slide_height
    _check_slide_size(width, height, parsed or TemplateConfig(), errors, warnings)
    layouts: dict[str, Any] = {}
    if parsed is not None:
        # With a config that could not be read, which layout a role gets is
        # unknown: showing the defaults would describe an export that will
        # not happen.
        layouts, layout_errors, layout_warnings = _resolve_layouts(prs, parsed)
        errors.extend(layout_errors)
        warnings.extend(layout_warnings)
    if dark_mode and parsed is not None and parsed.colors:
        warnings.append(TemplateIssue(
            "colors-ignored-in-dark-mode",
            "[colors] is not applied to a dark-mode export: dark mode uses the "
            "built-in dark palette, which is tuned for contrast on the dark slide.",
        ))
    report = TemplateValidationReport(
        template=template,
        config=config_source,
        slide_width_in=round(width / _EMU_PER_INCH, 3),
        slide_height_in=round(height / _EMU_PER_INCH, 3),
        available_layouts=tuple(layout.name for layout in prs.slide_layouts),
        layouts={role: layouts[role].name for role in ROLES if role in layouts},
        errors=tuple(errors),
        warnings=tuple(warnings),
    )
    return report, prs, layouts, parsed


def main(argv: list[str]) -> int:
    """CLI entry: ``validate-template [--config FILE] [--dark-mode] [--json] <template>``.

    Prints which layout each slide role would use and every error and
    warning. Exit code 0 when the template meets the contract, 2 when it does
    not, so a script can check a template before a long export.

    Example::

        thesisagents validate-template thesis.pptx --config thesis.toml
    """
    usage = "usage: validate-template [--config FILE] [--dark-mode] [--json] <template.pptx>"
    config: str | None = None
    dark_mode = as_json = False
    paths: list[str] = []
    index = 0
    while index < len(argv):
        argument = argv[index]
        if argument in ("-h", "--help"):
            print(usage)
            return 0
        if argument == "--config" and index + 1 < len(argv):
            config = argv[index + 1]
            index += 2
            continue
        if argument == "--dark-mode":
            dark_mode = True
        elif argument == "--json":
            as_json = True
        else:
            paths.append(argument)
        index += 1
    if len(paths) != 1:
        print(usage)
        return 2
    report = validate_template(paths[0], config, dark_mode=dark_mode)
    if as_json:
        print(json.dumps(report.to_dict(), ensure_ascii=False, indent=2))
    else:
        print(format_report(report))
    return 0 if report.ok else 2


def format_report(report: TemplateValidationReport) -> str:
    """The report as the lines a terminal shows.

    Example::

        Template thesis.pptx: OK
          slides: 13.33 x 7.50 in
          cover       -> Title Slide
          content     -> Title Only
          warning: [colors] is not applied to a dark-mode export: ...
    """
    lines = [f"Template {report.template}: {'OK' if report.ok else 'NOT USABLE'}"]
    if report.config:
        lines.append(f"  config: {report.config}")
    if report.slide_width_in is not None:
        lines.append(
            f"  slides: {report.slide_width_in:.2f} x {report.slide_height_in:.2f} in"
        )
    width = max((len(role) for role in report.layouts), default=0)
    lines.extend(
        f"  {role:<{width}} -> {report.layouts[role]}"
        for role in ROLES
        if role in report.layouts
    )
    lines.extend(f"  error: {issue.message}" for issue in report.errors)
    lines.extend(f"  warning: {issue.message}" for issue in report.warnings)
    return "\n".join(lines)


def _check_slide_size(
    width: int,
    height: int,
    config: TemplateConfig,
    errors: list[TemplateIssue],
    warnings: list[TemplateIssue],
) -> None:
    width_fits = abs(width - SLIDE_WIDTH) <= _SIZE_TOLERANCE
    height_fits = abs(height - SLIDE_HEIGHT) <= _SIZE_TOLERANCE
    if width_fits and height_fits:
        return
    size = f"{width / _EMU_PER_INCH:.2f} x {height / _EMU_PER_INCH:.2f} in"
    target = f"{SLIDE_WIDTH / _EMU_PER_INCH:.2f} x {SLIDE_HEIGHT / _EMU_PER_INCH:.2f} in (16:9)"
    if config.slide_size == SIZE_REJECT:
        errors.append(TemplateIssue(
            "slide-size",
            f"the template's slides are {size}, and the exporter lays out for {target}. "
            "Save a 16:9 version of the template (PowerPoint: Design > Slide Size), "
            f"or set slide_size = \"{SIZE_NORMALISE}\" to have the deck resized.",
        ))
        return
    warnings.append(TemplateIssue(
        "slide-size-normalised",
        f"the template's slides are {size}. The deck is resized to {target}, so "
        "artwork the template placed for its own size keeps its position and may "
        "sit off-centre. Save a 16:9 version of the template to avoid that, or set "
        f"slide_size = \"{SIZE_REJECT}\" to refuse such a template.",
    ))
