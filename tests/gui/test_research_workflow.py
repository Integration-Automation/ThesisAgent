"""GUI controls for citation snowballing, the literature library and deck templates.

Network is mocked: the search and the snowball are monkey-patched on the Search
page module. The library is a real SQLite file under ``tmp_path``.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from thesisagents.core.diagnostics import PaperRelation, RelationKind
from thesisagents.core.models import Paper, PaperCollection, Query
from thesisagents.core.snowball import DiscoveredPaper, SnowballResult
from thesisagents.gui.pages.deck import DeckPage
from thesisagents.gui.pages.search import SearchPage
from thesisagents.library import Library

SEARCH = "thesisagents.gui.pages.search"


def _paper(source_id: str, title: str, **fields) -> Paper:
    return Paper(
        source="arxiv", source_id=source_id, title=title, authors=("Author A",),
        year=fields.pop("year", 2024), venue=None,
        abstract=fields.pop("abstract", "An abstract about attention."),
        url=f"https://example.com/{source_id}", **fields,
    )


SEED = _paper("2401.00001", "Attention Mechanisms at Scale", arxiv_id="2401.00001")
FOUND = _paper("W9", "Sparse Attention Revisited", doi="10.1000/w9", year=2023)


def _collection(*papers: Paper) -> PaperCollection:
    return PaperCollection(Query("attention", ("arxiv",)), papers or (SEED,))


@pytest.fixture()
def page(qtbot, monkeypatch):
    """A Search page whose search returns ``SEED`` and counts its calls."""
    calls: list[str] = []

    async def fake_run_search(query, **_kwargs):
        calls.append(query.keywords)
        return _collection()

    async def fake_shutdown():
        return None

    monkeypatch.setattr(f"{SEARCH}.run_search", fake_run_search)
    monkeypatch.setattr(f"{SEARCH}.shutdown_clients", fake_shutdown)
    widget = SearchPage(ui_language="en")
    qtbot.addWidget(widget)
    widget.set_query_text("attention")
    widget.search_calls = calls
    return widget


def _titles(page: SearchPage) -> list[str]:
    return [paper.title for paper in page._collection.papers]  # noqa: SLF001


# ---------------------------------------------------------------- snowball


def test_the_snowball_selector_is_off_by_default_and_offers_three_directions(page):
    combo = page.snowball_combo()
    assert [combo.itemData(i) for i in range(combo.count())] == [
        None, "references", "cited_by", "both",
    ]
    assert combo.currentData() is None
    assert combo.itemText(0) == "Off"


def test_a_search_with_the_snowball_on_appends_what_it_finds(page, qtbot, monkeypatch):
    seen: dict[str, object] = {}

    async def fake_snowball(seeds, **kwargs):
        seen.update(kwargs, seeds=seeds)
        relation = PaperRelation(
            SEED.dedup_key(), FOUND.dedup_key(), RelationKind.CITED_BY, "openalex", 1
        )
        return SnowballResult(
            seeds=tuple(seeds), discovered=(DiscoveredPaper(FOUND, relation),),
            relations=(relation,),
        )

    monkeypatch.setattr(f"{SEARCH}.snowball", fake_snowball)
    page.snowball_combo().setCurrentIndex(3)                      # both
    page._on_search_clicked()  # noqa: SLF001

    qtbot.waitUntil(lambda: page.papers_model().rowCount() == 2, timeout=5000)
    assert _titles(page) == [SEED.title, FOUND.title]
    assert "Snowball: 1 new." in page.status_text()
    assert seen["direction"] == "both"
    assert seen["seeds"] == (SEED,)
    assert seen["known"] == (SEED,)
    assert seen["keywords"] == "attention"
    assert seen["max_total"] == 20                                # the CLI's default


def test_a_search_with_the_snowball_off_does_not_run_it(page, qtbot, monkeypatch):
    async def must_not_run(*_args, **_kwargs):
        raise AssertionError("the snowball ran")

    monkeypatch.setattr(f"{SEARCH}.snowball", must_not_run)
    page._on_search_clicked()  # noqa: SLF001

    qtbot.waitUntil(lambda: page.papers_model().rowCount() == 1, timeout=5000)
    assert "Snowball" not in page.status_text()


def test_the_snowball_summary_is_localised(qtbot):
    from thesisagents.gui.pages.search import _SearchOutcome

    widget = SearchPage(ui_language="zh-tw")
    qtbot.addWidget(widget)
    widget._on_search_finished(_SearchOutcome(_collection(), discovered=3))  # noqa: SLF001
    assert "滾雪球: 新增 3 篇。" in widget.status_text()


# ----------------------------------------------------------------- library


def test_ticking_add_stores_the_results_in_the_library(page, qtbot, tmp_path):
    library_path = tmp_path / "thesis.db"
    page.set_library_path(str(library_path))
    page.library_add_checkbox().setChecked(True)
    page._on_search_clicked()  # noqa: SLF001

    qtbot.waitUntil(lambda: "Library:" in page.status_text(), timeout=5000)
    assert "Library: 1 added, 0 already there." in page.status_text()
    with Library(library_path, create=False) as library:
        assert [entry.paper.title for entry in library.search("")] == [SEED.title]
        assert library.runs()[0].keywords == "attention"


def test_a_library_path_alone_does_not_store_the_results(page, qtbot, tmp_path):
    library_path = tmp_path / "thesis.db"
    page.set_library_path(str(library_path))
    page._on_search_clicked()  # noqa: SLF001

    qtbot.waitUntil(lambda: page.papers_model().rowCount() == 1, timeout=5000)
    assert "Library" not in page.status_text()
    assert not library_path.exists()


def test_add_without_a_library_file_stops_before_the_search(page):
    page.library_add_checkbox().setChecked(True)
    page._on_search_clicked()  # noqa: SLF001

    assert page.status_text() == "Enter a library file first."
    assert page.search_calls == []


def test_search_library_fills_the_table_without_searching(page, qtbot, tmp_path):
    library_path = tmp_path / "thesis.db"
    with Library(library_path) as library:
        off_topic = _paper("x1", "Protein Folding Notes", abstract="Structure prediction.")
        library.add_papers([SEED, FOUND, off_topic])
    page.set_library_path(str(library_path))
    page._on_library_search_clicked()  # noqa: SLF001

    qtbot.waitUntil(lambda: page.papers_model().rowCount() == 2, timeout=5000)
    assert set(_titles(page)) == {SEED.title, FOUND.title}        # the two on attention
    assert page.status_text() == "2 of 3 library papers."
    assert page.search_calls == []                                # no network search
    assert page._export_button.isEnabled()  # noqa: SLF001


def test_search_library_with_an_empty_query_lists_the_library(page, qtbot, tmp_path):
    library_path = tmp_path / "thesis.db"
    with Library(library_path) as library:
        library.add_papers([SEED, FOUND])
    page.set_library_path(str(library_path))
    page.set_query_text("")
    page._on_library_search_clicked()  # noqa: SLF001

    qtbot.waitUntil(lambda: page.papers_model().rowCount() == 2, timeout=5000)
    assert page.status_text() == "2 of 2 library papers."


def test_search_library_on_a_missing_file_is_an_error_and_creates_nothing(page, qtbot, tmp_path):
    missing = tmp_path / "typo.db"
    page.set_library_path(str(missing))
    page._on_library_search_clicked()  # noqa: SLF001

    qtbot.waitUntil(lambda: "no library at" in page.status_text(), timeout=5000)
    assert not missing.exists()
    assert page._library_search_button.isEnabled()  # noqa: SLF001


def test_search_library_needs_a_library_file(page):
    page._on_library_search_clicked()  # noqa: SLF001
    assert page.status_text() == "Enter a library file first."


def test_export_keeps_the_identifier_checks_in_the_library(
    page, qtbot, monkeypatch, tmp_path
):
    captured: dict[str, object] = {}

    def fake_export(collection, options, *, verification_cache=None):
        captured["cache"] = verification_cache
        return {"bib": Path(options.out_dir) / "fake.bib"}

    monkeypatch.setattr(f"{SEARCH}.export_collection", fake_export)
    monkeypatch.setattr(
        f"{SEARCH}.QFileDialog.getExistingDirectory", lambda *_a, **_k: str(tmp_path)
    )
    page._on_search_finished(_collection())  # noqa: SLF001
    page.set_library_path(str(tmp_path / "thesis.db"))
    page._on_export_clicked()  # noqa: SLF001

    qtbot.waitUntil(lambda: "fake.bib" in page.status_text(), timeout=5000)
    assert type(captured["cache"]).__name__ == "LibraryVerificationCache"


# --------------------------------------------------------------- templates


def _deck_page(qtbot, tmp_path) -> DeckPage:
    widget = DeckPage(ui_language="en")
    qtbot.addWidget(widget)
    widget._out_dir_input.setText(str(tmp_path / "out"))  # noqa: SLF001
    widget.set_collection(_collection())
    return widget


def _template(path: Path) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    prs.save(str(path))
    return path


def test_the_template_fields_reach_the_export(qtbot, monkeypatch, tmp_path):
    captured: dict[str, object] = {}

    def fake_export(collection, options):
        captured["options"] = options
        return {"pptx": Path(options.out_dir) / "fake.pptx"}

    monkeypatch.setattr("thesisagents.gui.pages.deck.export_collection", fake_export)
    widget = _deck_page(qtbot, tmp_path)
    widget.set_template_path("  thesis.pptx ")
    widget.set_template_config_path("thesis.toml")
    widget._on_export_clicked()  # noqa: SLF001

    qtbot.waitUntil(lambda: "Wrote" in widget.status_text(), timeout=3000)
    options = captured["options"]
    assert (options.pptx_template, options.pptx_template_config) == ("thesis.pptx", "thesis.toml")


def test_empty_template_fields_mean_the_built_in_design(qtbot, monkeypatch, tmp_path):
    captured: dict[str, object] = {}

    def fake_export(collection, options):
        captured["options"] = options
        return {"pptx": Path(options.out_dir) / "fake.pptx"}

    monkeypatch.setattr("thesisagents.gui.pages.deck.export_collection", fake_export)
    widget = _deck_page(qtbot, tmp_path)
    widget._on_export_clicked()  # noqa: SLF001

    qtbot.waitUntil(lambda: "Wrote" in widget.status_text(), timeout=3000)
    options = captured["options"]
    assert (options.pptx_template, options.pptx_template_config) == (None, None)


def test_a_template_config_without_a_template_is_reported_at_once(qtbot, monkeypatch, tmp_path):
    def must_not_run(*_args, **_kwargs):
        raise AssertionError("the export ran")

    monkeypatch.setattr("thesisagents.gui.pages.deck.export_collection", must_not_run)
    widget = _deck_page(qtbot, tmp_path)
    widget.set_template_config_path("thesis.toml")
    widget._on_export_clicked()  # noqa: SLF001

    assert "needs pptx_template as well" in widget.status_text()
    assert widget._export_button.isEnabled() is True  # noqa: SLF001


def test_a_template_that_does_not_meet_the_contract_is_shown_and_nothing_is_written(
    qtbot, tmp_path
):
    """No monkeypatch: the real export runs and the template check stops it."""
    config = tmp_path / "thesis.toml"
    config.write_text('[layouts]\ntable = "Data"\n', encoding="utf-8")
    widget = _deck_page(qtbot, tmp_path)
    widget.verify_identifiers_checkbox().setChecked(False)
    widget.set_template_path(str(_template(tmp_path / "thesis.pptx")))
    widget.set_template_config_path(str(config))
    widget._on_export_clicked()  # noqa: SLF001

    qtbot.waitUntil(
        lambda: "does not meet the template contract" in widget.status_text(), timeout=5000
    )
    assert '[layouts] table = "Data"' in widget.status_text()
    assert not list((tmp_path / "out").glob("*.pptx"))


def test_a_valid_template_builds_the_deck(qtbot, tmp_path):
    widget = _deck_page(qtbot, tmp_path)
    widget.verify_identifiers_checkbox().setChecked(False)
    widget.set_template_path(str(_template(tmp_path / "thesis.pptx")))
    widget._on_export_clicked()  # noqa: SLF001

    qtbot.waitUntil(lambda: "Wrote" in widget.status_text(), timeout=10000)
    deck = Presentation(str(next((tmp_path / "out").glob("*.pptx"))))
    assert {slide.slide_layout.name for slide in deck.slides} == {"Blank"}


# ------------------------------------------- library as the deck tab's cache


def test_the_deck_export_keeps_identifier_checks_in_the_search_tabs_library(
    qtbot, monkeypatch, tmp_path
):
    captured: dict[str, object] = {}

    def fake_export(collection, options, *, verification_cache=None):
        captured["cache"] = verification_cache
        return {"pptx": Path(options.out_dir) / "fake.pptx"}

    monkeypatch.setattr("thesisagents.gui.pages.deck.export_collection", fake_export)
    widget = _deck_page(qtbot, tmp_path)
    widget.set_library_path(f"  {tmp_path / 'thesis.db'} ")
    widget._on_export_clicked()  # noqa: SLF001

    qtbot.waitUntil(lambda: "Wrote" in widget.status_text(), timeout=3000)
    assert type(captured["cache"]).__name__ == "LibraryVerificationCache"


def test_the_main_window_passes_the_library_path_to_the_deck_tab(qtbot):
    from thesisagents.gui.main_window import MainWindow

    window = MainWindow(ui_language="en")
    qtbot.addWidget(window)
    window.search_page().set_library_path(" thesis.db ")
    assert window.deck_page()._library_path == "thesis.db"  # noqa: SLF001
    window.search_page().set_library_path("")
    assert window.deck_page()._library_path == ""  # noqa: SLF001
