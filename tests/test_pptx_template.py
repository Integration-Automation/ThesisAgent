"""The PPTX template contract: validation, layouts by role, style overrides."""

from __future__ import annotations

import asyncio
import json
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from thesisagents import cli as cli_module
from thesisagents import mcp as mcp_pkg
from thesisagents.core.models import (
    ExportOptions,
    Paper,
    PaperCollection,
    PaperSummary,
    Query,
)
from thesisagents.exporters import export_collection, pptx, pptx_edit, template
from thesisagents.exporters.audit import audit_deck
from thesisagents.exporters.review import review_deck
from thesisagents.exporters.template import (
    ROLES,
    TemplateConfig,
    TemplateError,
    load_template_config,
    parse_template_config,
    validate_template,
)
from thesisagents.mcp.server import ToolError

NAVY = (0x1F, 0x3A, 0x66)
WHITE = (0xFF, 0xFF, 0xFF)
GREEN = (0x0B, 0x3D, 0x2E)


# ------------------------------------------------------------------ helpers


def _make_template(
    path: Path,
    *,
    size: tuple[float, float] = (13.333, 7.5),
    rename: dict[str, str] | None = None,
    keep: tuple[str, ...] | None = None,
    title_box: tuple[float, float] | None = (0.3, 0.9),
) -> Path:
    """A template derived from python-pptx's own, saved to ``path``.

    ``rename`` renames layouts, ``keep`` drops every layout not listed (by its
    original name), ``title_box`` is the ``(top, height)`` in inches given to
    the title placeholder of "Title Only" (``None`` leaves the stock 1.55 in
    bottom edge, which is below the title area).
    """
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(size[0]), Inches(size[1])
    for layout in list(prs.slide_layouts):
        if keep is not None and layout.name not in keep:
            prs.slide_layouts.remove(layout)
            continue
        if layout.name == "Title Only" and title_box is not None:
            title = layout.placeholders[0]
            title.top, title.height = Inches(title_box[0]), Inches(title_box[1])
        if rename and layout.name in rename:
            layout.name = rename[layout.name]
    prs.save(str(path))
    return path


def _config(tmp_path: Path, text: str, name: str = "template.toml") -> Path:
    path = tmp_path / name
    path.write_text(text, encoding="utf-8")
    return path


def _collection(papers: int = 1) -> PaperCollection:
    summary = PaperSummary(
        language="en",
        research_question="Can a template carry the deck?",
        pain_points=(("Templates drift", ("Layouts are found by index", "Fonts are fixed")),),
        headline_metrics=(("Layouts mapped", "6", ""),),
        technique_table=(("Role mapping", "Picks a layout per slide role"),),
        contributions_detailed=(("A contract", "Named layouts replace a fixed index."),),
    )
    items = tuple(
        Paper(
            source="arxiv", source_id=f"2401.0000{n}", title=f"Template Paper {n}",
            authors=("Ada Lovelace",), year=2024, venue="NeurIPS",
            abstract="An abstract about templates.", url=f"https://arxiv.org/abs/2401.0000{n}",
            summary=summary,
        )
        for n in range(1, papers + 1)
    )
    return PaperCollection(Query("templates", ("arxiv",)), items)


def _export(tmp_path: Path, collection=None, **options) -> Presentation:
    options.setdefault("verify_identifiers", False)
    written = export_collection(
        collection or _collection(),
        ExportOptions(formats=("pptx",), out_dir=str(tmp_path / "out"), **options),
    )
    return Presentation(str(written["pptx"]))


def _shape(slide, name: str):
    return next((shape for shape in slide.shapes if shape.name == name), None)


def _run_colours(shape) -> set[tuple[int, int, int]]:
    return {
        tuple(run.font.color.rgb)
        for paragraph in shape.text_frame.paragraphs
        for run in paragraph.runs
        if run.text.strip()
    }


def _codes(issues) -> list[str]:
    return [issue.code for issue in issues]


# ------------------------------------------------------- built-in unchanged


def test_the_built_in_deck_is_unchanged(tmp_path):
    prs = _export(tmp_path)
    slides = list(prs.slides)

    assert (prs.slide_width, prs.slide_height) == (pptx._SLIDE_WIDTH, pptx._SLIDE_HEIGHT)
    assert {slide.slide_layout.name for slide in slides} == {"Blank"}
    assert all(not list(slide.placeholders) for slide in slides)
    assert _shape(slides[0], "accent_left") is not None           # navy cover panel
    assert _run_colours(_shape(slides[0], "title")) == {WHITE}
    content = slides[1]
    assert tuple(_shape(content, "accent_top").fill.fore_color.rgb) == NAVY
    assert _run_colours(_shape(content, "title")) == {WHITE}
    assert _shape(content, "title").text_frame.paragraphs[0].runs[0].font.name == "Inter"


def test_the_shared_title_area_constant_matches_the_exporter():
    # The contract's limit for a title placeholder is where the exporter puts
    # the line under the title, and the canvas is the exporter's canvas.
    assert pptx.TITLE_AREA_BOTTOM is template.TITLE_AREA_BOTTOM
    assert (template.SLIDE_WIDTH, template.SLIDE_HEIGHT) == (pptx._SLIDE_WIDTH, pptx._SLIDE_HEIGHT)
    assert Inches(1.24) < template.TITLE_AREA_BOTTOM <= pptx._BODY_TOP


# --------------------------------------------------------------- validation


def test_a_stock_16_9_template_is_valid_and_every_role_falls_back_to_blank(tmp_path):
    report = validate_template(_make_template(tmp_path / "t.pptx"))

    assert report.ok and report.errors == () and report.warnings == ()
    assert (report.slide_width_in, report.slide_height_in) == (13.333, 7.5)
    assert report.layouts == dict.fromkeys(ROLES, "Blank")
    assert "Title Only" in report.available_layouts


def test_layouts_named_after_a_role_are_picked_up_without_a_config(tmp_path):
    path = _make_template(
        tmp_path / "t.pptx",
        rename={"Title Slide": "Cover", "Title Only": "content", "Section Header": "Q&A"},
    )
    layouts = validate_template(path).layouts

    assert layouts["cover"] == "Cover"                 # case-insensitive
    assert layouts["content"] == "content"
    assert layouts["qa"] == "Q&A"                      # alias of the qa role
    assert layouts["table"] == layouts["references"] == "content"   # fall back to content


def test_a_config_maps_roles_to_layout_names(tmp_path):
    path = _make_template(tmp_path / "t.pptx")
    config = _config(tmp_path, '[layouts]\ncover = "Title Slide"\ntable = "title only"\n')
    layouts = validate_template(path, config).layouts

    assert layouts["cover"] == "Title Slide"
    assert layouts["table"] == "Title Only"
    assert layouts["content"] == "Blank"


def test_a_layout_name_that_does_not_exist_is_an_actionable_error(tmp_path):
    path = _make_template(tmp_path / "t.pptx")
    report = validate_template(path, _config(tmp_path, '[layouts]\ntable = "Data"\n'))

    assert not report.ok
    assert _codes(report.errors) == ["layout-not-found"]
    message = report.errors[0].message
    assert '[layouts] table = "Data"' in message
    assert "'Title Only'" in message and "'Blank'" in message      # what there is


def test_a_template_with_no_layout_for_content_is_refused(tmp_path):
    path = _make_template(tmp_path / "t.pptx", keep=("Title Slide", "Title and Content"))
    report = validate_template(path)

    assert _codes(report.errors) == ["no-content-layout"]
    assert "Add a layout named 'content' or 'Blank'" in report.errors[0].message
    assert '[layouts] content = "..."' in report.errors[0].message


def test_a_missing_required_title_placeholder_is_an_actionable_error(tmp_path):
    path = _make_template(tmp_path / "t.pptx")
    config = _config(tmp_path, '[placeholders]\ntitle = ["content"]\n')    # content -> Blank
    report = validate_template(path, config)

    assert _codes(report.errors) == ["missing-title-placeholder"]
    message = report.errors[0].message
    assert "layout 'Blank' has no title placeholder" in message
    assert "View > Slide Master" in message
    assert "remove 'content' from [placeholders] title" in message


def test_a_title_placeholder_that_reaches_into_the_content_is_refused(tmp_path):
    path = _make_template(tmp_path / "t.pptx", title_box=None)             # ends at 1.55 in
    config = _config(
        tmp_path, '[layouts]\ncontent = "Title Only"\n[placeholders]\ntitle = ["content"]\n'
    )
    report = validate_template(path, config)

    assert _codes(report.errors) == ["title-placeholder-overlaps-content"]
    assert "ends 1.55 in from the top" in report.errors[0].message
    assert "at or above 1.40 in" in report.errors[0].message


def test_a_non_16_9_template_is_normalised_with_a_warning_by_default(tmp_path):
    path = _make_template(tmp_path / "t.pptx", size=(10, 7.5))
    report = validate_template(path)

    assert report.ok
    assert _codes(report.warnings) == ["slide-size-normalised"]
    assert "10.00 x 7.50 in" in report.warnings[0].message
    assert "13.33 x 7.50 in (16:9)" in report.warnings[0].message
    prs = _export(tmp_path, pptx_template=str(path))
    assert (prs.slide_width, prs.slide_height) == (pptx._SLIDE_WIDTH, pptx._SLIDE_HEIGHT)


def test_a_non_16_9_template_is_rejected_when_the_config_says_so(tmp_path):
    path = _make_template(tmp_path / "t.pptx", size=(10, 7.5))
    report = validate_template(path, _config(tmp_path, 'slide_size = "reject"\n'))

    assert _codes(report.errors) == ["slide-size"]
    assert "Design > Slide Size" in report.errors[0].message
    assert 'slide_size = "normalise"' in report.errors[0].message


def test_a_file_that_is_not_a_presentation_is_reported_not_raised(tmp_path):
    notes = tmp_path / "notes.pptx"
    notes.write_text("not a presentation", encoding="utf-8")
    report = validate_template(notes)

    assert _codes(report.errors) == ["unreadable-template"]
    assert "cannot open" in report.errors[0].message
    assert _codes(validate_template(tmp_path / "missing.pptx").errors) == ["unreadable-template"]


def test_a_broken_config_is_reported_with_the_template_findings(tmp_path):
    path = _make_template(tmp_path / "t.pptx", size=(10, 7.5))
    report = validate_template(path, _config(tmp_path, "[colors]\nprimary = 5\n"))

    assert _codes(report.errors) == ["invalid-config"]
    assert "[colors] primary must be a colour" in report.errors[0].message
    assert _codes(report.warnings) == ["slide-size-normalised"]


def test_the_report_serialises(tmp_path):
    path = _make_template(tmp_path / "t.pptx")
    payload = validate_template(path, _config(tmp_path, '[layouts]\ntable = "Data"\n')).to_dict()

    assert payload["ok"] is False
    assert payload["errors"][0]["code"] == "layout-not-found"
    assert payload["layouts"]["content"] == "Blank"
    assert json.loads(json.dumps(payload)) == payload


# ------------------------------------------------------------------- config


def test_a_config_is_read_from_toml_or_json(tmp_path):
    toml = _config(
        tmp_path,
        'slide_size = "reject"\n'
        '[layouts]\ncover = "Title Slide"\n'
        '[placeholders]\ntitle = ["content", "table"]\n'
        '[fonts]\nlatin = "Source Sans 3"\neast_asian = "Noto Sans TC"\n'
        '[colors]\nprimary = "#0B3D2E"\nhighlight = "0E7C61"\n'
        "[chrome]\nheader_band = false\n",
    )
    as_json = _config(
        tmp_path,
        json.dumps({
            "slide_size": "reject",
            "layouts": {"cover": "Title Slide"},
            "placeholders": {"title": ["content", "table"]},
            "fonts": {"latin": "Source Sans 3", "east_asian": "Noto Sans TC"},
            "colors": {"primary": "#0B3D2E", "highlight": "0E7C61"},
            "chrome": {"header_band": False},
        }),
        name="template.json",
    )
    for path in (toml, as_json):
        config = load_template_config(path)
        assert config.slide_size == "reject"
        assert config.layouts == {"cover": "Title Slide"}
        assert config.title_placeholder_roles == {"content", "table"}
        assert (config.font_latin, config.font_east_asian) == ("Source Sans 3", "Noto Sans TC")
        assert config.colors == {"primary": GREEN, "highlight": (0x0E, 0x7C, 0x61)}
        assert (config.header_band, config.cover_panel) == (False, True)
        assert config.source == str(path)


def test_an_empty_config_changes_nothing():
    assert parse_template_config({}) == TemplateConfig()


@pytest.mark.parametrize(
    ("data", "fragment"),
    [
        ({"margins": {"left": 1}}, "[margins] is not supported: margins are fixed"),
        ({"sizes": {"title": 40}}, "[sizes] is not supported: font sizes are fixed"),
        ({"theme": {}}, "unknown setting 'theme'"),
        ({"slide_size": "stretch"}, 'slide_size must be "normalise" or "reject"'),
        ({"layouts": {"appendix": "X"}}, "[layouts] has no setting 'appendix'"),
        ({"layouts": {"cover": ""}}, "[layouts] cover must be a layout name"),
        ({"layouts": "Title"}, "[layouts] must be a table of settings"),
        ({"placeholders": {"title": ["cover"]}}, "[placeholders] title cannot include 'cover'"),
        ({"placeholders": {"title": "content"}}, "[placeholders] title must be a list of roles"),
        ({"fonts": {"latin": ""}}, "[fonts] latin must be a font family name"),
        ({"fonts": {"heading": "X"}}, "[fonts] has no setting 'heading'"),
        ({"colors": {"primary": "navy"}}, "[colors] primary must be a colour like"),
        ({"colors": {"primary": "#FFF3B0"}}, "is too light"),
        ({"colors": {"muted": "#C0392B"}}, "the red this project bans as a text colour"),
        ({"colors": {"accent": "#000000"}}, "[colors] has no setting 'accent'"),
        ({"chrome": {"header_band": "no"}}, "[chrome] header_band must be true or false"),
    ],
)
def test_a_config_problem_names_the_setting(data, fragment):
    with pytest.raises(TemplateError) as raised:
        parse_template_config(data)
    assert fragment in str(raised.value)


def test_every_config_problem_is_reported_at_once():
    with pytest.raises(TemplateError) as raised:
        parse_template_config(
            {"margins": {}, "fonts": {"latin": ""}, "colors": {"primary": "navy"}},
            source="thesis.toml",
        )
    message = str(raised.value)
    assert "template config thesis.toml has 3 problem(s)" in message
    assert message.count("\n  - ") == 3


def test_a_light_primary_says_why_it_must_be_dark():
    with pytest.raises(TemplateError, match="the fill behind white titles"):
        parse_template_config({"colors": {"primary": "#FDE68A"}})


@pytest.mark.parametrize(
    ("name", "text", "fragment"),
    [
        ("c.yaml", "a: 1", "must be a .toml or .json file"),
        ("c.toml", "[layouts", "is not valid TOML"),
        ("c.json", "{", "is not valid JSON"),
        ("c.json", "[1, 2]", "must hold a table of settings"),
    ],
)
def test_an_unreadable_config_says_what_is_wrong(tmp_path, name, text, fragment):
    with pytest.raises(TemplateError) as raised:
        load_template_config(_config(tmp_path, text, name=name))
    assert fragment in str(raised.value)


def test_a_missing_config_file_is_a_template_error(tmp_path):
    with pytest.raises(TemplateError, match="cannot read template config"):
        load_template_config(tmp_path / "nope.toml")


# ---------------------------------------------------------------- rendering


@pytest.fixture()
def mapped(tmp_path):
    """A template with a layout per role and the config that maps them."""
    path = _make_template(tmp_path / "thesis.pptx")
    config = _config(
        tmp_path,
        "[layouts]\n"
        'cover = "Title Slide"\n'
        'section = "Section Header"\n'
        'content = "Title Only"\n'
        'table = "Title and Content"\n'
        'references = "Two Content"\n'
        'qa = "Picture with Caption"\n',
    )
    return str(path), str(config)


def test_a_valid_custom_template_renders_each_role_on_its_layout(tmp_path, mapped):
    path, config = mapped
    prs = _export(tmp_path, _collection(papers=2), pptx_template=path, pptx_template_config=config)
    slides = list(prs.slides)
    by_layout: dict[str, list] = {}
    for slide in slides:
        by_layout.setdefault(slide.slide_layout.name, []).append(slide)

    assert slides[0].slide_layout.name == "Title Slide"            # cover
    assert len(by_layout["Section Header"]) == 2                   # one divider per paper
    assert len(by_layout["Picture with Caption"]) == 2             # one Q&A per paper
    assert len(by_layout["Two Content"]) == 1                      # references
    assert all(_shape(slide, "title") for slide in by_layout["Title and Content"])   # tables
    table_slides = by_layout["Title and Content"]
    assert any(shape.has_table for slide in table_slides for shape in slide.shapes)
    assert "Title Only" in by_layout                               # everything else
    assert "Blank" not in by_layout


def test_inherited_placeholders_are_removed_from_every_slide(tmp_path, mapped):
    path, config = mapped
    prs = _export(tmp_path, pptx_template=path, pptx_template_config=config)
    # No "Click to add title" box ships: the exporter drew its own shapes.
    assert all(not list(slide.placeholders) for slide in prs.slides)
    assert all(_shape(slide, "title") is not None for slide in prs.slides)


def test_a_template_deck_round_trips_and_passes_the_deck_audits(tmp_path, mapped):
    path, config = mapped
    written = export_collection(
        _collection(),
        ExportOptions(
            formats=("pptx",), out_dir=str(tmp_path / "out"), verify_identifiers=False,
            pptx_template=path, pptx_template_config=config,
        ),
    )
    reopened = Presentation(str(written["pptx"]))                  # opens with python-pptx
    assert len(reopened.slides) > 5
    review = review_deck(written["pptx"])
    assert list(review.overflow) == []
    assert review.hard_contrast == []


def test_the_title_goes_into_the_templates_title_placeholder_when_asked(tmp_path):
    path = _make_template(tmp_path / "t.pptx", title_box=(0.25, 0.8))
    config = _config(
        tmp_path,
        '[layouts]\ncontent = "Title Only"\n[placeholders]\ntitle = ["content", "table"]\n',
    )
    prs = _export(tmp_path, pptx_template=str(path), pptx_template_config=str(config))
    # Every role falls back to the content layout here, so the cover is on
    # "Title Only" too. It is told apart by how its title was made.
    in_placeholder = [s for s in prs.slides if _shape(s, "title").is_placeholder]
    assert len(in_placeholder) >= 3
    slide = in_placeholder[0]
    title = _shape(slide, "title")

    assert slide.slide_layout.name == "Title Only"
    assert (title.top, title.height) == (Inches(0.25), Inches(0.8))     # the template's box
    assert title.has_text_frame and title.text_frame.text.strip()
    assert _run_colours(title) == {NAVY}                                # explicit, not white
    assert title.text_frame.paragraphs[0].runs[0].font.size == pptx.Pt(pptx._SECTION_TITLE_PT)
    assert _shape(slide, "accent_top") is None                  # the template owns the header
    assert [p for p in slide.placeholders] == [title]                   # and nothing else is left
    # The cover is not a title-placeholder role: it keeps its own text box.
    assert not _shape(prs.slides[0], "title").is_placeholder


def test_a_title_placeholder_deck_has_no_invisible_text_in_dark_mode(tmp_path):
    path = _make_template(tmp_path / "t.pptx", title_box=(0.25, 0.8))
    config = _config(
        tmp_path, '[layouts]\ncontent = "Title Only"\n[placeholders]\ntitle = ["content"]\n'
    )
    written = export_collection(
        _collection(),
        ExportOptions(
            formats=("pptx",), out_dir=str(tmp_path / "out"), verify_identifiers=False,
            pptx_template=str(path), pptx_template_config=str(config), dark_mode=True,
        ),
    )
    assert [issue for issue in audit_deck(written["pptx"]) if issue.hard] == []


def test_custom_fonts_are_applied_to_every_run(tmp_path):
    path = _make_template(tmp_path / "t.pptx")
    config = _config(tmp_path, '[fonts]\nlatin = "Source Sans 3"\neast_asian = "Noto Sans TC"\n')
    prs = _export(tmp_path, pptx_template=str(path), pptx_template_config=str(config))
    runs = [
        run
        for slide in prs.slides
        for shape in slide.shapes if shape.has_text_frame
        for paragraph in shape.text_frame.paragraphs
        for run in paragraph.runs
    ]

    assert runs
    assert {run.font.name for run in runs} == {"Source Sans 3"}
    east_asian = {
        run._r.get_or_add_rPr().find(pptx.qn("a:ea")).get("typeface") for run in runs
    }
    assert east_asian == {"Noto Sans TC"}


def test_a_latin_only_font_override_keeps_the_languages_east_asian_font(tmp_path):
    path = _make_template(tmp_path / "t.pptx")
    config = _config(tmp_path, '[fonts]\nlatin = "Source Sans 3"\n')
    prs = _export(
        tmp_path, pptx_template=str(path), pptx_template_config=str(config), language="zh-tw"
    )
    run = _shape(prs.slides[0], "title").text_frame.paragraphs[0].runs[0]

    assert run.font.name == "Source Sans 3"
    assert run._r.get_or_add_rPr().find(pptx.qn("a:ea")).get("typeface") == "Microsoft JhengHei UI"


def test_custom_colours_replace_the_palette_as_text_and_as_fill(tmp_path):
    path = _make_template(tmp_path / "t.pptx")
    config = _config(tmp_path, '[colors]\nprimary = "#0B3D2E"\nhighlight = "#0E7C61"\n')
    prs = _export(tmp_path, pptx_template=str(path), pptx_template_config=str(config))
    text_colours: set = set()
    fills: set = set()
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_table:
                fills.add(tuple(shape.table.cell(0, 0).fill.fore_color.rgb))      # header row
            elif shape.name in ("accent_top", "accent_left", "accent_rule"):
                fills.add(tuple(shape.fill.fore_color.rgb))
            if shape.has_text_frame:
                text_colours |= _run_colours(shape)

    assert GREEN in text_colours and (0x0E, 0x7C, 0x61) in text_colours
    assert NAVY not in text_colours and (0x25, 0x63, 0xEB) not in text_colours
    assert fills == {GREEN, (0x0E, 0x7C, 0x61)}             # band, cover, table header, rule
    assert WHITE in text_colours                                    # titles stay white on the fill


def test_custom_colours_are_not_applied_in_dark_mode_and_that_is_reported(tmp_path):
    path = _make_template(tmp_path / "t.pptx")
    config = _config(tmp_path, '[colors]\nprimary = "#0B3D2E"\n')
    report = validate_template(path, config, dark_mode=True)
    prs = _export(
        tmp_path, pptx_template=str(path), pptx_template_config=str(config), dark_mode=True
    )

    assert report.ok and _codes(report.warnings) == ["colors-ignored-in-dark-mode"]
    assert _codes(validate_template(path, config).warnings) == []
    colours = {
        colour
        for slide in prs.slides
        for shape in slide.shapes if shape.has_text_frame
        for colour in _run_colours(shape)
    }
    assert GREEN not in colours and NAVY not in colours             # the dark palette won


def test_switching_the_chrome_off_leaves_readable_text(tmp_path):
    path = _make_template(tmp_path / "t.pptx")
    config = _config(tmp_path, "[chrome]\nheader_band = false\ncover_panel = false\n")
    prs = _export(tmp_path, pptx_template=str(path), pptx_template_config=str(config))
    slides = list(prs.slides)

    for slide in slides:
        for name in ("accent_top", "accent_rule", "accent_left"):
            assert _shape(slide, name) is None
    # White-on-navy text would now be white on the template's light background.
    assert _run_colours(_shape(slides[0], "title")) == {NAVY}
    assert _run_colours(_shape(slides[0], "meta")) == {(0x55, 0x55, 0x55)}
    assert _run_colours(_shape(slides[1], "title")) == {NAVY}
    every = {c for slide in slides for s in slide.shapes if s.has_text_frame and not s.has_table
             for c in _run_colours(s)}
    assert WHITE not in every


def test_only_the_cover_panel_can_be_switched_off(tmp_path):
    path = _make_template(tmp_path / "t.pptx")
    config = _config(tmp_path, "[chrome]\ncover_panel = false\n")
    prs = _export(tmp_path, pptx_template=str(path), pptx_template_config=str(config))

    assert _shape(prs.slides[0], "accent_left") is None
    assert _shape(prs.slides[1], "accent_top") is not None
    assert _run_colours(_shape(prs.slides[1], "title")) == {WHITE}


# ---------------------------------------------------------- export preflight


def test_export_stops_before_writing_when_the_template_does_not_meet_the_contract(tmp_path):
    path = _make_template(tmp_path / "t.pptx", size=(10, 7.5))
    config = _config(tmp_path, 'slide_size = "reject"\n[layouts]\ntable = "Data"\n')

    with pytest.raises(TemplateError) as raised:
        _export(tmp_path, pptx_template=str(path), pptx_template_config=str(config))
    message = str(raised.value)
    assert message.startswith("[pptx] template ")
    assert "does not meet the template contract" in message
    assert "Design > Slide Size" in message and '[layouts] table = "Data"' in message
    assert not list((tmp_path / "out").glob("*.pptx"))


def test_a_template_with_fewer_than_seven_layouts_no_longer_crashes(tmp_path):
    # slide_layouts[6] used to be taken whatever the template held.
    path = _make_template(tmp_path / "t.pptx", keep=("Title Slide", "Blank"))
    prs = _export(tmp_path, pptx_template=str(path))
    assert {slide.slide_layout.name for slide in prs.slides} == {"Blank"}


def test_a_config_without_a_template_is_refused():
    with pytest.raises(ValueError, match="needs pptx_template as well"):
        ExportOptions(formats=("pptx",), out_dir="x", pptx_template_config="c.toml")


def test_a_slide_can_be_added_to_a_template_deck_with_few_layouts(tmp_path):
    path = _make_template(tmp_path / "t.pptx", keep=("Title Only", "Blank"))
    written = export_collection(
        _collection(),
        ExportOptions(
            formats=("pptx",), out_dir=str(tmp_path / "out"), verify_identifiers=False,
            pptx_template=str(path),
        ),
    )
    before = len(Presentation(str(written["pptx"])).slides)
    pptx_edit.add_slide(written["pptx"], title="Appendix", body="One more point")
    prs = Presentation(str(written["pptx"]))

    assert len(prs.slides) == before + 1
    added = prs.slides[-1]
    assert added.slide_layout.name == "Blank"
    assert not list(added.placeholders)
    assert _shape(added, "title").text_frame.text == "Appendix"


# ---------------------------------------------------------------------- CLI


@pytest.fixture()
def stub_search(monkeypatch):
    calls: list[str] = []

    async def fake_run_search(query, **_kwargs):  # NOSONAR async stub
        calls.append(query.keywords)
        return _collection()

    async def fake_shutdown():  # NOSONAR async stub
        return None

    monkeypatch.setattr(cli_module, "run_search", fake_run_search)
    monkeypatch.setattr(cli_module, "shutdown_clients", fake_shutdown)
    return calls


def _cli(tmp_path, *extra: str) -> int:
    return cli_module.main(
        ["--query", "templates", "--source", "arxiv", "--no-pdf", "--export", "pptx",
         "--no-verify-identifiers", "--out", str(tmp_path / "cli-out"), *extra]
    )


def test_cli_builds_the_deck_on_the_template(tmp_path, stub_search, mapped):
    path, config = mapped
    assert _cli(tmp_path, "--pptx-template", path, "--pptx-template-config", config) == 0

    deck = Presentation(str(next((tmp_path / "cli-out").glob("*.pptx"))))
    assert deck.slides[0].slide_layout.name == "Title Slide"
    assert "Title Only" in {slide.slide_layout.name for slide in deck.slides}


def test_cli_checks_the_template_before_the_search(tmp_path, stub_search, capsys):
    path = _make_template(tmp_path / "t.pptx")
    config = _config(tmp_path, '[layouts]\ntable = "Data"\n')

    code = _cli(tmp_path, "--pptx-template", str(path), "--pptx-template-config", str(config))

    assert code == 2
    assert stub_search == []                                        # the search never ran
    err = capsys.readouterr().err
    assert "does not meet the template contract" in err
    assert '[layouts] table = "Data"' in err


def test_cli_prints_template_warnings_even_when_quiet(tmp_path, stub_search, capsys):
    path = _make_template(tmp_path / "t.pptx", size=(10, 7.5))
    assert _cli(tmp_path, "--pptx-template", str(path), "--quiet") == 0
    assert "Template warning: the template's slides are 10.00 x 7.50 in" in capsys.readouterr().err


def test_cli_template_config_needs_a_template(tmp_path, stub_search):
    with pytest.raises(SystemExit, match="--pptx-template-config needs --pptx-template FILE"):
        _cli(tmp_path, "--pptx-template-config", "c.toml")
    assert stub_search == []


def test_validate_template_subcommand_reports_a_usable_template(tmp_path, mapped, capsys):
    path, config = mapped
    assert cli_module.main(["validate-template", path, "--config", config]) == 0

    out = capsys.readouterr().out
    assert f"Template {path}: OK" in out
    assert "slides: 13.33 x 7.50 in" in out
    assert "cover      -> Title Slide" in out
    assert "references -> Two Content" in out


def test_validate_template_subcommand_exits_two_and_lists_the_errors(tmp_path, capsys):
    path = _make_template(tmp_path / "t.pptx", size=(10, 7.5))
    config = _config(tmp_path, 'slide_size = "reject"\n')

    assert cli_module.main(["validate-template", str(path), "--config", str(config)]) == 2
    out = capsys.readouterr().out
    assert "NOT USABLE" in out
    assert "  error: the template's slides are 10.00 x 7.50 in" in out


def test_validate_template_subcommand_can_emit_json(tmp_path, capsys):
    path = _make_template(tmp_path / "t.pptx")
    assert cli_module.main(["validate-template", "--json", str(path)]) == 0
    payload = json.loads(capsys.readouterr().out)
    assert payload["ok"] is True and payload["layouts"]["qa"] == "Blank"


def test_validate_template_subcommand_needs_one_template(capsys):
    assert cli_module.main(["validate-template"]) == 2
    assert "usage: validate-template" in capsys.readouterr().out


# ---------------------------------------------------------------------- MCP


def _call(name: str, **kwargs):
    result = asyncio.run(mcp_pkg.build_server().call_tool(name, kwargs))
    if isinstance(result, tuple):
        result = result[0]
    result = getattr(result, "content", result)
    return json.loads(next(block.text for block in result if getattr(block, "text", None)))


def test_mcp_lists_eighteen_tools_including_the_template_check():
    names = [tool.name for tool in asyncio.run(mcp_pkg.build_server().list_tools())]
    assert "pptx_validate_template" in names
    assert len(names) == 18


def test_mcp_validate_template_returns_the_report(tmp_path):
    path = _make_template(tmp_path / "t.pptx")
    config = _config(tmp_path, '[layouts]\ntable = "Data"\n')

    usable = _call("pptx_validate_template", path=str(path))
    broken = _call("pptx_validate_template", path=str(path), config=str(config))

    assert usable["ok"] is True and usable["layouts"]["content"] == "Blank"
    assert broken["ok"] is False and broken["errors"][0]["code"] == "layout-not-found"


def test_mcp_export_builds_on_a_template_or_fails_with_the_report(tmp_path, mapped):
    path, config = mapped
    paper = _collection().papers[0].to_dict()
    arguments = {
        "papers": [paper], "keywords": "templates", "formats": ["pptx"],
        "out_dir": str(tmp_path / "mcp-out"), "verify_identifiers": False,
    }

    written = _call("export", **arguments, pptx_template=path, pptx_template_config=config)
    deck = Presentation(written["pptx_path"])
    assert deck.slides[0].slide_layout.name == "Title Slide"

    bad = _config(tmp_path, '[layouts]\ntable = "Data"\n', name="bad.toml")
    with pytest.raises(ToolError, match="does not meet the template contract"):
        _call("export", **arguments, pptx_template=path, pptx_template_config=str(bad))
