"""Degree-thesis ORAL-DEFENCE deck (學位論文口試簡報) for 陳冠穎's master's thesis
《基於大語言模型和思維鏈推理的程式碼審查框架》, built via the `thesis-deck-author`
flow.

Source of truth
---------------
Hand-authored from the candidate's OWN thesis (read in full):
  D:\\Codes\\Code-Review-Framework-Combining-Large-Language-Models-and-Chain-of-Thought-Reasoning
  - paper/論文_v2.9.docx        ... the authoritative final manuscript (defended)
  - paper/REWRITE_BRIEF.md       ... frozen-facts + integrated-framework thesis axis
  - codes/ + datas/ + prthinker/ ... numbers verified against real source / score files

Every number on these slides is the candidate's own — none invented.  Where the
repo `score.md` (legacy run) and the v2.9 manuscript diverged, the deck follows
the **v2.9 manuscript** (the document actually being defended): CRSCORE++
1.00 / 0.79 / 0.86 vs baseline 0.67 / 0.57 / 0.63, base model gemma-4-31B-it with
the prior Qwen3-Coder-30B-A3B scoring comparably.  Honest scoping is preserved:
IDE integration = MCP + pre-commit only (no editor plugin); the 17 research-level
mechanisms are implemented + unit-tested but NOT end-to-end evaluated; global
rules are 7 static + a conditional 8th; no cross-backend comparison numbers.

White + blue academic style (dark_mode=False) per the user's request.

The candidate supplied the final bilingual thesis title for this build, and
asked for 標楷體 (CJK) + Times New Roman (Latin) typography — both applied by
the ``_restyle_deck`` post-export pass below (the exporter's zh-tw default is
Inter + Microsoft JhengHei UI, which this deck overrides).

Run from the project root:  .venv/Scripts/python.exe scripts/regen_chen2026_codereview.py
"""
from __future__ import annotations

import sys
from pathlib import Path

_PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(_PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(_PROJECT_ROOT))

from pptx import Presentation  # noqa: E402
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402

from thesisagents.core.models import (  # noqa: E402
    ExportOptions,
    Paper,
    PaperCollection,
    PaperSummary,
    Query,
    RqResult,
)
from thesisagents.exporters import overflow as _overflow  # noqa: E402
from thesisagents.exporters.pptx import (  # noqa: E402
    _BODY_WIDTH,
    _BRAND_GREY,
    _DARK_BODY_TEXT,
    _MARGIN_X,
    PptxExporter,
    _add_textbox,
    _set_east_asian_typeface,
)

_EMU_PER_INCH = 914400

OUT_DIR = "exports"
FILENAME_STEM = "chen2026-llm-cot-codereview-zh-tw"
LANGUAGE = "zh-tw"

# Final bilingual thesis title supplied by the candidate for this build. The
# Chinese title is the cover headline (Paper.title); the English title is
# surfaced on the cover's reserved subtitle slot by _restyle_deck.
CHINESE_TITLE = "基於多階段思維鏈與檢索增強生成之大語言模型程式碼審查框架設計與實作"
ENGLISH_TITLE = (
    "Design and Implementation of a Large Language Model-Based Code Review "
    "Framework Using Multi-Stage Chain-of-Thought Reasoning and "
    "Retrieval-Augmented Generation"
)

# Candidate-requested typography for this deck. PowerPoint consults the <a:ea>
# (East-Asian) slot for CJK code points and the <a:latin> slot for Latin code
# points within the SAME run, so a mixed 中文+English run renders each script in
# its correct face without splitting runs.
_CJK_FONT = "標楷體"          # DFKai-SB, the Traditional-Chinese 楷書 face
_LATIN_FONT = "Times New Roman"

# Real thesis flow diagrams (圖二 訓練流程 / 圖三 程式碼審查流程), copied from the
# candidate's own paper/_media19 so this script is self-contained.
_FIG = _PROJECT_ROOT / "assets" / "figures" / "chen2026"


def _build_summary() -> PaperSummary:
    return PaperSummary(
        language=LANGUAGE,
        # ===== 1. 緒論:1.2 動機 — LLM 程式碼審查的三大問題(全篇主軸)=========
        pain_points=(
            (
                "人工審查是品質把關的瓶頸",
                (
                    "審查可在整測前攔下 50–70% 缺陷",
                    "但耗時耗人力、標準不一",
                    "直接用 LLM(大型語言模型)審查帶出三大問題 ↓",
                ),
            ),
            (
                "問題一:幻覺,審查意見不可信",
                (
                    "編造看似合理卻不存在的呼叫",
                    "錯誤審查比沒有更危險",
                    "需可查證依據,而非自由發揮",
                ),
            ),
            (
                "問題二:輸出不穩定,難以複現",
                (
                    "同一變更給出不一致結論",
                    "結果難以複現與比較",
                    "需拆解任務、固定流程與規則",
                ),
            ),
            (
                "問題三:缺乏領域規範,脫離脈絡",
                (
                    "不懂你專案的命名與安全規範",
                    "建議與專案脈絡脫節",
                    "需即時注入規則,而非重訓模型",
                ),
            ),
        ),
        # ===== 1.3 研究問題(RQ callout;在此首次解釋 CRSCORE++)============
        research_question=(
            "如何讓 LLM 程式碼審查同時具備一致性、可解釋性與可維護性,"
            "正面回應幻覺、輸出不穩定與缺乏領域規範三大問題,"
            "並在 CRSCORE++(程式碼審查品質的 LLM 自動評分基準)等指標上優於既有方法?"
        ),
        # ===== 摘要 / 4.4 主要量化成果(KPI;採 v2.9 manuscript 數字)========
        headline_metrics=(
            ("CRSCORE++ 完整性", "1.00", "基準 0.67 · 近滿分,幾乎不漏點"),
            ("CRSCORE++ 相關性", "0.86", "基準 0.63 · 評論貼合變更"),
            ("CRSCORE++ 簡潔性", "0.79", "基準 0.57 · 冗言更少"),
            ("基準測試資料", "44 筆", "Python · 三型 · 人工驗證"),
        ),
        # ===== 2.x 關鍵技術 → 角色(術語在「主要貢獻」頁首次解釋,此處為角色對照)===
        technique_table=(
            ("思維鏈 CoT", "把審查拆成多步,先逐步推理再下結論"),
            ("RAG 檢索增強", "生成前先檢索專案規則並注入提示詞,降低幻覺"),
            ("知識蒸餾 KD", "讓學生模型照教師的範例與推理軌跡學會審查"),
            ("QLoRA", "在有限顯示記憶體下微調,訓練量化、推論不量化"),
            ("FAISS", "Facebook AI 向量索引庫,高速近似最近鄰檢索"),
            ("LLM-as-a-Judge", "以獨立的 LLM 當評審,對審查意見逐維度評分"),
        ),
        # ===== 2.3 文獻比較與研究缺口(表一精簡;header 為首列)============
        literature_table=(
            ("方法 / 研究", "多階段", "微調", "RAG", "評估方式", "主要限制"),
            ("CRScore [7]", "✗", "✗", "✗", "3 維 1–5 分", "評分維度粗"),
            ("CRScore++ [10]", "✗", "✗", "✗", "3 維 + RL", "缺深度語義"),
            ("LLaMA-Reviewer [16]", "✗", "✓", "✗", "BLEU / ROUGE", "無語義層次"),
            ("LAURA [20]", "✗", "✓", "✓", "字串重疊", "未拆解推理"),
            ("本研究 Ours", "✓", "✓", "✓", "5 維百分制 + 人工", "詳見 §6.3"),
        ),
        # ===== 3.1 / 3.4 系統與審查流程(端到端,最易懂的一條線)===========
        system_flow=(
            "GitHub PR(Pull Request,合併程式碼的審查申請)觸發 → CI 呼叫 prthinker",
            "RAG 檢索層:把程式碼變更向量化,以餘弦相似度 ≥ 0.7 取最相關的專案規則",
            "五步思維鏈逐步審查:摘要 → 初步審查 → Linter(靜態語法檢查) → 程式碼異味 → 總結,每步落 .md",
            "全域規則統一前綴:七條靜態準則 + 命中時動態追加的第八條 RAG 規則",
            "JudgeStep 保守聚合 → APPROVE / REQUEST_CHANGES / COMMENT,回貼 PR",
        ),
        # ===== 3.1 / 3.2 / 3.4 論文原圖(取自 論文_v2.9.docx 之最新高解析版)=====
        figures=(
            (
                "系統架構:從 GitHub PR 到審查回貼的八層設計",
                str(_FIG / "system_architecture.png"),
                (
                    "進入與 CI → 邊緣伺服(FastAPI 非同步)→ 思維鏈審查管線",
                    "知識檢索(RAG / FAISS / Repo 知識圖譜)→ 後端抽象(可替換推論後端)",
                    "推論與離線訓練 → 狀態快取 → 輸出與監控",
                ),
            ),
            (
                "訓練流程:把教師審查能力蒸餾進可部署的學生模型",
                str(_FIG / "training_flow.png"),
                (
                    "教師模型蒸餾出帶推理軌跡資料 → 建 JSONL → 前處理",
                    "QLoRA 載入學生基座 → 掛上並訓練 LoRA 適配器",
                    "預設不合併、輸出適配器 → 部署(bf16、雙卡)→ 開機守門擋 OOM 配置",
                ),
            ),
            # 審查流程(論文 圖三)為很高的單欄流程圖,縮到投影片高度後寬度過窄而難辨識;
            # 其五步流程已以「系統架構」頁的文字管線(system_flow)清楚呈現,故此處不重複嵌圖。
        ),
        # ===== 3.2-3.4 方法細節(每節先白話直覺,再技術;2 節/頁 → 2 頁)=====
        method_sections=(
            (
                "多階段 CoT:把一次大審查拆成五個單一子任務",
                (
                    "直覺:與其要模型一口氣審完,不如分五步、每步只做一件事",
                    "順序固定:摘要 → 初步審查 → Linter → 程式碼異味 → 總結",
                    "每步產物落 <step>_result.md,中途不失、可追溯 → 回應輸出不穩定",
                ),
            ),
            (
                "全域規則:七條靜態準則 + 條件式第八條",
                (
                    "build_global_rule_template 為每步注入統一的規則前綴",
                    "七條:可讀一致 / 命名 / 軟工原則 / 邏輯正確 / 效能安全 / 文件測試 / 評分風格",
                    "當且僅當 RAG 檢索到非空專案規則,才在七條後動態追加第八條(非固定八條)",
                ),
            ),
            (
                "RAG 領域規則:把專案規範臨時塞進提示詞",
                (
                    "直覺:模型不懂你專案的規範,審查前臨時檢索最相關幾條規則注入,不必重訓",
                    "嵌入向量 → FAISS IndexFlatIP(L2 正規化使內積即餘弦),匯入時建索引一次",
                    "餘弦 ≥ 0.7 命中者注入,另以 repo 符號知識圖譜接地,抑制幻覺函式名",
                ),
            ),
            (
                "KD + QLoRA:把教師審查能力輕量化移轉到可部署的學生",
                (
                    "直覺:讓學生模型照教師寫好的範例與推理軌跡學會審查(知識蒸餾)",
                    "以少樣本 CoT 提示向教師模型蒸餾出帶推理軌跡的訓練資料",
                    "QLoRA 4-bit NF4 微調學生(LoRA r=64),訓練量化、推論用 bf16、適配器可拆卸",
                ),
            ),
        ),
        # ===== 4.2 評估方法(三層 + 公正性;2 節/頁 → 1 頁)================
        evaluation_sections=(
            (
                "三層評估:自動 + 自研 + 人工互補",
                (
                    "CRSCORE++:三維各 1–5 分(完整性 / 簡潔性 / 相關性),沿用既有基準",
                    "自研 Our:五維百分制(可讀性 / 可維護性 / 正確性 / 跨評論覆蓋 / 完整性),解析度更細",
                    "人工評估:同五維,把關權限控管、併發競爭、資源生命週期等隱性問題",
                ),
            ),
            (
                "公正性設計:評審獨立於受測模型",
                (
                    "LLM-as-a-Judge 以 GPT-5 為主評審、Gemini-3 第二評審交叉",
                    "評審模型獨立於被審查的學生模型,避免自評偏差",
                    "自動評分再以人工交叉驗證,尤其是可維護性與正確性",
                ),
            ),
        ),
        # ===== 4.1 基準測試資料組成(44 筆,逐型筆數經程式碼核實)===========
        paper_tables=(
            (
                "基準測試資料組成(44 筆)",
                (
                    ("資料型態", "筆數", "說明"),
                    ("bad_data", "4", "刻意埋入缺陷的程式"),
                    ("code_diff", "12", "PR 變更差異"),
                    ("only_code", "28", "單純程式片段"),
                    ("合計", "44", "GPT-5 與 Copilot 生成 · 人工驗證 · Python"),
                ),
                (
                    "小於 20 筆難以覆蓋多維錯誤型態,上百筆又增加人工標註成本,取 44 筆為平衡",
                    "涵蓋語法 / 命名 / 設計模式 / 錯誤型態等多維問題",
                ),
            ),
        ),
        # ===== 4 / 1.3 研究問題清單 =========================================
        research_questions=(
            ("RQ1", "多階段 CoT + 微調框架,相較單一提示詞與 CRSCORE++,能否在完整性 / 簡潔性 / 相關性顯著提升審查品質?"),
            ("RQ2", "固定參數規模下,僅加入多階段提示詞、不微調,是否即可帶來品質提升?(分離提示詞與微調)"),
            ("RQ3", "多階段提示詞與微調各自的邊際貢獻為何?何者為品質改善的主因?"),
            ("RQ4", "自動 LLM-as-a-Judge 的結論,能否由獨立人工評分在可維護性與正確性上交叉驗證?"),
        ),
        # ===== 4.4 / 4.5 各 RQ 結果(v2.9 manuscript 數字,逐字沿用)=========
        rq_results=(
            RqResult(
                rq_id="RQ1",
                question="整合框架在 CRSCORE++ 三維是否全面優於既有基準?",
                table=(
                    ("CRSCORE++ 維度", "本框架 Ours", "CRSCORE++ 基準"),
                    ("完整性 comprehensiveness", "1.00", "0.67"),
                    ("相關性 relevance", "0.86", "0.63"),
                    ("簡潔性 conciseness", "0.79", "0.57"),
                ),
                analysis=(
                    "本框架(gemma-4-31B-it,與前代 Qwen3-Coder-30B-A3B 相當)三維皆高於基準",
                    "完整性近滿分,代表該提的問題幾乎不漏,相關性 0.86 代表評論貼合變更",
                    "意義:審查涵蓋度明顯高於基準(1.00 對 0.67),更少漏看真正該提的問題",
                ),
            ),
            RqResult(
                rq_id="RQ3",
                question="多階段提示詞與微調,各自的邊際貢獻為何?",
                table=(
                    ("維度(GPT-5 評分)", "單一提示詞", "多階段 Ours"),
                    ("可維護性 Maintainability", "85", "94"),
                    ("正確性 Correctness", "82", "90"),
                    ("簡潔性 Conciseness", "78", "96"),
                    ("跨評論覆蓋 Coverage", "90", "97"),
                ),
                analysis=(
                    "多階段提示詞流程是整合框架內單項貢獻最大的一環,折算約 +34 分",
                    "LoRA 微調邊際約 +2 分,但解決不同問題:有限資源下保留並穩定教師能力",
                    "重點:+34 與 +2 是貢獻分解、互補,而非『流程勝過微調、可砍掉微調』",
                ),
            ),
            RqResult(
                rq_id="RQ4",
                question="自動評分的結論,能否由獨立人工評分交叉驗證?",
                table=(
                    ("維度(人工評分)", "基礎模型", "多階段 Ours"),
                    ("可維護性 Maintainability", "79.88", "86.25"),
                    ("正確性 Correctness", "80.75", "87.75"),
                    ("跨評論覆蓋 Coverage", "74.13", "86.25"),
                    ("簡潔性 Conciseness", "76.25", "84.25"),
                ),
                analysis=(
                    "人工在可維護性、正確性、跨評論覆蓋三維與自動評分方向一致 → 交叉驗證成立",
                    "唯可讀性出現自動(92)高於人工(83.5)的系統性偏差,與文獻 [11] LLM 評審語感類較寬鬆一致",
                    "意義:機器打的分數人工大致認同,尤其在『改得對不對、好不好維護』上",
                ),
            ),
        ),
        # ===== 1.5 研究貢獻(cap 4;前三為已驗證核心 + 評估方法)===========
        contributions_detailed=(
            (
                "一、多階段 CoT 審查流程",
                "思維鏈(Chain-of-Thought, CoT)把審查拆成五步、每步落檔,並注入七條全域"
                "規則,消融顯示為框架內單項貢獻最大者(約 +34 分)。",
            ),
            (
                "二、RAG 領域規則注入",
                "RAG(檢索增強生成)以 FAISS(向量索引庫)餘弦 ≥ 0.7 把最相關專案規則"
                "動態注入為條件式第八條,回應缺乏領域規範。",
            ),
            (
                "三、KD + QLoRA 輕量化",
                "知識蒸餾(KD)+ QLoRA(4-bit 量化 LoRA 微調)把教師大模型的審查能力移轉到"
                "30B 級可部署的學生模型(Qwen3-Coder-30B-A3B / gemma-4-31B-it),適配器可拆卸。",
            ),
            (
                "四、LLM-as-a-Judge-Our 五維評估",
                "以另一個 LLM 當評審(LLM-as-a-Judge)做自研百分制五維評分,並由 GPT-5、"
                "Gemini-3 與人工三方交叉驗證,評審獨立於受測模型。",
            ),
        ),
        # ===== 5.1 核心觀察(整合框架主軸;callout)========================
        core_observation=(
            "核心發現:並非以單一元件取代其餘,而是整合多階段 CoT(一致性、可追溯)、"
            "RAG 條件式規則(可維護、領域貼合)、KD+QLoRA(有限資源落地)三者,才能讓 LLM "
            "程式碼審查同時具備一致性、可解釋性與可維護性,正面回應幻覺、輸出不穩定與缺乏領域"
            "規範三大問題。消融中多階段流程貢獻最大(+34)、微調互補(+2),缺一不可。"
        ),
        # ===== 5.3 研究限制 =================================================
        limitations=(
            "資料規模:44 筆 Python,尚未驗證 C++ / Java / Go 與其他專案類型的泛化",
            "評審偏差:GPT-5、Gemini-3 同屬商用 LLM,可能共享分布",
            "微調範圍:僅單一學生配置,基底換 gemma-4-31B-it 後評審評分尚待完整重評",
            "部署實證:尚未於真實開發團隊長期試行",
            "統計檢定:44 筆未做顯著性檢定,RAG 規則注入未單獨消融",
        ),
        # ===== 5.4 未來工作 =================================================
        future_work=(
            "跨後端品質 / 成本 / 延遲評估(gemma 重放已完成質性檢核,量化重評為下一步)",
            "以 PR 作者反饋累積學習語料(dismissed / accepted),評估長期成效",
            "跨平台(GitLab / Bitbucket / Gitea)與多模型仲裁",
            "隨附框架十七項研究級擴充機制的端到端實證(目前已實作 + 單元測試,效益未評估)",
        ),
        model="hand-authored:thesis-deck-author",
    )


def _build_paper() -> Paper:
    return Paper(
        source="local",
        source_id="chen2026-llm-cot-codereview",
        title=CHINESE_TITLE,
        authors=("陳冠穎",),
        year=2026,
        venue="國立高雄師範大學 · 軟體工程與管理學系 · 碩士學位論文 · 指導教授:李文廷 博士",
        abstract="",
        url="",
        doi=None,  # 學位論文無出版 DOI → 觸發 own-thesis 路徑(不產生來源頁 / 參考文獻頁)
        summary=_build_summary(),
    )


def _bake_font_scale(text_frame, scale: float) -> None:
    """Persist an explicit ``fontScale`` on a shrink-to-fit text frame.

    The exporter sets ``auto_size = TEXT_TO_FIT_SHAPE``, which python-pptx
    writes as a bare ``<a:normAutofit/>`` with no ``fontScale`` attribute.
    LibreOffice computes the shrink on the fly, but PowerPoint renders such a
    box at FULL font size (it only recomputes the scale when the text / box is
    edited) — so long titles overrun their box and overlap the content below.
    Baking the scale (percent × 1000, e.g. 78% → ``"78000"``) makes every
    renderer, PowerPoint included, draw the text already-shrunk.
    """
    pct = max(1, min(100000, int(round(scale * 100000))))
    body_pr = text_frame._txBody.find(qn("a:bodyPr"))  # noqa: SLF001
    norm = body_pr.find(qn("a:normAutofit")) if body_pr is not None else None
    if norm is None:  # not a shrink-to-fit frame — nothing to bake
        return
    norm.set("fontScale", str(pct))


def _restyle_deck(out_path: str) -> None:
    """Post-export pass for this deck's candidate-specific requests + fit fixes.

    1. Surface the English thesis title on the cover. The exporter renders one
       cover title (the Chinese ``Paper.title``); the cover's reserved subtitle
       slot (top 3.3") is empty for a Chinese-only title, so the English title
       goes there beneath the Chinese headline.
    2. Trim the per-slide running subtitle to the thesis title alone. The
       exporter's ``_add_paper_subtitle`` packs ``title · venue`` into a 0.28"
       one-line slot; with the long final title that line wraps to two and
       overruns into the body. The cover already carries the venue / author, so
       the running header only needs the title.
    3. Override every run's typography to 標楷體 (CJK) + Times New Roman (Latin).
       The exporter's ``_apply_typography`` pins zh-tw to Inter + Microsoft
       JhengHei UI; we re-walk every run AFTER export and rewrite both the
       ``<a:latin>`` (Times New Roman) and ``<a:ea>`` (標楷體) typeface slots so
       PowerPoint renders each script in its requested face.
    4. Bake a ``fontScale`` into any shrink-to-fit box that still overruns at
       full size (e.g. the RQ-question callout), so PowerPoint draws it shrunk
       rather than spilling over — see ``_bake_font_scale``.

    Why a post-pass and not exporter changes: the font choice, the second cover
    title, and the running-header trim are this deck's alone, not a project-wide
    default — re-running this script reproduces the deck exactly, with no edit
    to the shared exporter.
    """
    prs = Presentation(out_path)

    cover = prs.slides[0]
    _add_textbox(
        cover, name="subtitle", text=ENGLISH_TITLE,
        left=_MARGIN_X, top=Inches(3.3),
        width=_BODY_WIDTH, height=Inches(1.3),
        font_pt=18, colour=_DARK_BODY_TEXT,
        align=PP_ALIGN.CENTER, shrink_to_fit=True,
    )

    # (2) Shorten the running subtitle on every content slide to the title
    # only. Rewriting the text drops the run's size / colour, so restore them.
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame or shape.name != "paper_subtitle":
                continue
            tf = shape.text_frame
            runs = [r for p in tf.paragraphs for r in p.runs]
            size = runs[0].font.size if runs else Pt(14)
            tf.text = CHINESE_TITLE
            run = tf.paragraphs[0].runs[0]
            run.font.size = size or Pt(14)
            run.font.color.rgb = _BRAND_GREY

    # (3) Typography override on every run.
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.name = _LATIN_FONT
                    _set_east_asian_typeface(run, _CJK_FONT)

    # (4) Bake fontScale where a shrink-to-fit box still overruns at full size.
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            tf = shape.text_frame
            if tf.auto_size != MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE:
                continue
            box_h = (shape.height or 0) / _EMU_PER_INCH
            est = _overflow._text_height_in(tf, shape.width or 0)  # noqa: SLF001
            if box_h > 0 and est > box_h * 0.97:
                _bake_font_scale(tf, max(0.55, (box_h * 0.95) / est))

    prs.save(out_path)


def main() -> None:
    collection = PaperCollection(
        query=Query(keywords="LLM CoT code review framework", sources=("local",)),
        papers=(_build_paper(),),
    )
    options = ExportOptions(
        formats=("pptx",),
        out_dir=OUT_DIR,
        filename_stem=FILENAME_STEM,
        language=LANGUAGE,
        # 口試簡報採白色 + 藍學術風(白底、深藍標題 / 內文、藍色強調)。
        dark_mode=False,
        # 先完整渲染所有作者內容,再依張數與時長修剪(避免 cap 靜默截斷)。
        max_slides_per_paper=30,
    )
    out_path = PptxExporter().export(collection, options)
    _restyle_deck(out_path)
    print(f"saved: {out_path}")


if __name__ == "__main__":
    main()
